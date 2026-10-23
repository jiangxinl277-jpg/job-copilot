#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Parse a candidate resume (docx / pdf / txt) into a structured profile.

Output: JSON with text, sections, keyword counts and capability tags.
Used by batch_score.py so the pipeline can accept *any* resume, not just the
hand-maintained profile.json.

Usage:
  python3 scripts/parse_resume.py --input resume.docx --out data/resume_profile.json
"""

import argparse
import json
import os
import re
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "scripts"))

from match_evidence import KW2TAG  # noqa: E402

SECTION_PATTERNS = [
    ("education", r"教育|学历|education"),
    ("experience", r"实习|工作经历|职业经历|experience|internship"),
    ("projects", r"项目|作品|project"),
    ("skills", r"技能|证书|语言|skills|certificat"),
    ("honors", r"荣誉|奖项|honor|award"),
]


def read_text(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        from docx import Document
        d = Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.extend(c.text for c in row.cells)
        return "\n".join(parts)
    if ext == ".pdf":
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            return "\n".join((p.extract_text() or "") for p in pdf.pages)
    return open(path, encoding="utf-8", errors="ignore").read()


def split_sections(text):
    sections = {"header": [], "education": [], "experience": [], "projects": [], "skills": [], "honors": []}
    current = "header"
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        head = None
        for name, pat in SECTION_PATTERNS:
            # pat 里有 | 交替，必须用非捕获组包起来，否则 fullmatch 会被拆开
            if len(line) <= 24 and re.fullmatch(r"[^\n]{0,12}(?:" + pat + r")[^\n]{0,12}", line, flags=re.I):
                head = name
                break
        if head:
            current = head
            continue
        sections[current].append(line)
    return sections


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True)
    ap.add_argument("--out", default=os.path.join(ROOT, "data", "resume_profile.json"))
    a = ap.parse_args()

    text = read_text(a.input)
    sections = split_sections(text)
    keywords = {k: text.count(k) for k in KW2TAG if text.count(k)}
    tags = set()
    for k in keywords:
        tags.update(KW2TAG.get(k, []))

    profile = {
        "source": os.path.abspath(a.input),
        "chars": len(text),
        "text": text,
        "sections": sections,
        "keywords": dict(sorted(keywords.items(), key=lambda kv: -kv[1])),
        "tags": sorted(tags),
        "lines": [l for l in (x.strip() for x in text.splitlines()) if l],
    }
    os.makedirs(os.path.dirname(os.path.abspath(a.out)) or ".", exist_ok=True)
    open(a.out, "w", encoding="utf-8").write(json.dumps(profile, ensure_ascii=False, indent=1))

    print(json.dumps({
        "chars": profile["chars"],
        "sections": {k: len(v) for k, v in sections.items()},
        "keywords": list(profile["keywords"])[:12],
        "tags": profile["tags"],
        "out": a.out,
    }, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
